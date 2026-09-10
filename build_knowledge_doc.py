# -*- coding: utf-8 -*-
"""生成《课题背景知识文档_五场景与仿真原理.docx》
内容：①五个泛化场景说明文字（论文 Table 6/7 定义 + 系统装备库编排）
      ②本课题仿真原理与可靠性论证（基于 engine.py 代码实证）
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
for sec in doc.sections:
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.8)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def title(text, size=18):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    p.paragraph_format.space_after = Pt(6)
    return p


def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(15)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p


def h3(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    return p


def body(text, bold=False, italic=False, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.first_line_indent = Pt(22)
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(11)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return p


def bullet(text, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Pt(22 + level * 18)
    p.paragraph_format.first_line_indent = Pt(-11)
    r = p.add_run(("▪ " if level == 0 else "– ") + text)
    r.font.size = Pt(11)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return p


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    return p


def make_table(headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htxt in enumerate(headers):
        cell = t.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(htxt)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "D9E2F3")
        cell._tc.get_or_add_tcPr().append(shd)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j > 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            r.font.size = Pt(9.5)
            r.font.name = "Times New Roman"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    return t


# ============================ 封面标题 ============================
title("课题背景知识文档")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("反无人机行动方案智能生成系统 —— 泛化场景说明与仿真原理")
r.bold = True
r.font.size = Pt(14)
r.font.name = "Times New Roman"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("（供导师审阅 · 背景知识汇总）")
r.font.size = Pt(10.5)
r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
p.paragraph_format.space_after = Pt(16)

# ============================ 第一部分：五个场景 ============================
h1("第一部分　五个泛化场景说明")

body("本系统在跨场景泛化实验中定义了五个场景族（coastal joint assault 沿海联合突击、urban hub defense 城市枢纽防御、"
     "mountain corridor reconnaissance 山地走廊侦察、river crossing breakthrough 渡河突破、island resupply corridor 岛屿补给走廊保障），"
     "覆盖进攻（assault）、防御（defense）、侦察（recon）、保障（sustain）四类任务类型。场景定义、地形天气、兵力数量与威胁系数来自论文 "
     "Table 6；我方/敌方装备构成为基于系统装备库（用户指导手册 7.3 节）的说明性编排，供理解场景全貌使用。")
note("说明：文中加粗字段为论文/系统的官方定义（数据可核查）；装备构成与威胁系数解读为说明性编排，如需精确引用可标注“示例构成”。")

# ---- 场景1 ----
h2("场景一　沿海联合突击（Coastal Joint Assault）")
h3("任务")
body("夺取登陆窗口，压制敌方岸防与防空节点，为后续两栖兵力上陆和扩展创造条件——即“先扫清滩头威胁，再送人上岸”。"
     "这是五个场景中对抗强度最高、要素最全的场景，也是论文主消融实验使用的场景。", bold=False)
h3("地形与天气")
body("海岸-城市复合地形（coastal-urban），降雨（rain）。雨雾降低光学探测效能，对态势感知（awareness）形成压制。")
h3("兵力对比与敌我装备")
body("我方 5 个作战单元 vs 敌方 3 个作战单元（全场景最大规模）。")
bullet("敌方（说明性编排）：岸防火力节点、近岸防空导弹阵地、侦察与自杀式无人机蜂群，以及水面快艇袭扰——典型“海岸线多层防御”。")
bullet("我方（示例编排）：JY-17B 低空搜索雷达（态势感知）、陆盾-3000 弹炮合一防空系统（火力+防护）、天猎-1 拦截无人机（机动拦截）、"
       "LW-30 激光反无人机系统（低成本精确毁伤）、DWL-200 车载电子对抗（电磁压制）。")
h3("威胁系数解读")
body("威胁度 0.74（全场景最高）、电磁对抗强度 0.68、时间压力 0.77（全场景最高）——即“敌强、电磁环境差、时间紧”，"
     "方案必须同时处理好火力、电子战与节奏三件事。")
h3("涉及能力维度")
body("火力、机动、态势感知、电子对抗、指挥控制——五个域全开，最能体现系统自适应能力（对应实验中任务成功率增益最大 +5.15 点）。")

# ---- 场景2 ----
h2("场景二　城市枢纽防御（Urban Hub Defense）")
h3("任务")
body("守住交通与通信枢纽，抵御敌方装甲部队的城市突击——典型“城市要点防御”，重点是防突破而非主动出击。")
h3("地形与天气")
body("城市（urban），浓雾（fog）。建筑物遮蔽加雾天，光电探测大打折扣，且存在平民问题（civilian presence）约束。")
h3("兵力对比与敌我装备")
body("我方 4 vs 敌方 2。")
bullet("敌方（说明性编排）：装甲突击分队（主攻轴线）、伴随小型无人机侦察、隐蔽的火箭筒/迫击炮火力点。")
bullet("我方（示例编排）：无线电侦测站（被动侦察，抗雾抗遮挡）、JY-17B 低空雷达、红旗-17A 近程防空导弹（保卫枢纽上空）、"
       "CHL-802 便携式干扰枪（近距离软杀伤、控制附带损伤）。")
h3("威胁系数解读")
body("威胁度 0.69、电磁对抗 0.54、时间压力 0.66——中等强度；核心难点是“城市三维空间 + 平民保护 + 软杀伤优先”的约束组合，"
     "而不是硬火力。")
h3("涉及能力维度")
body("态势感知、电子对抗、防护、指挥控制——城市环境对“避免误伤”和“精确识别”要求极高。")
h3("实验备注")
body("此场景任务成功率增益 +1.04（五场景最小）——因为纯 LLM 基线本身已较高（68.05），可解释为“基线强、提升空间小”。")

# ---- 场景3 ----
h2("场景三　山地走廊侦察（Mountain Corridor Reconnaissance）")
h3("任务")
body("在低能见度条件下收集敌方火力节点与补给线情报——侦察而非交战，方案重点是“看得远、藏得住、传得回”。")
h3("地形与天气")
body("山地走廊（mountain corridor），多云（cloudy）。山地遮挡严重，无线电视距受限。")
h3("兵力对比与敌我装备")
body("我方 3 vs 敌方 2（全场景最小规模，轻型化编成）。")
bullet("敌方（说明性编排）：隐蔽火炮/火箭炮火力节点、补给车队、少量无人机监视。")
bullet("我方（示例编排）：侦察无人机（抵近侦察）、DWL-500 光电探测系统（光学侦察）、无线电侦测站（被动定位敌方通信与雷达信号）。")
h3("威胁系数解读")
body("威胁度 0.58、电磁对抗 0.46、时间压力 0.55——全场景最低威胁组；难点不在“打不打得过”，而在“看不看得见、连不连得通”"
     "（山地对 awareness 与 C2 的考验）。")
h3("涉及能力维度")
body("态势感知、机动、指挥控制——侦察类场景的典型维度子集。")
h3("实验备注")
body("任务成功率增益 +1.86，纯 LLM 基线（68.59）本就较高，Full 提升到 70.45。")

# ---- 场景4 ----
h2("场景四　渡河突破（River Crossing Breakthrough）")
h3("任务")
body("在敌方有防御的情况下开辟渡河窗口并巩固桥头堡——进攻性场景，“过河”和“站稳”两件事都要做到，"
     "天然考验“突破（breach）”环节。")
h3("地形与天气")
body("河岸平原（river-plain），阴天（overcast）。")
h3("兵力对比与敌我装备")
body("我方 4 vs 敌方 2。")
bullet("敌方（说明性编排）：对岸防御阵地（反坦克/防空火力）、渡口区雷场与障碍、无人机监视与袭扰。")
bullet("我方（示例编排）：天猎-1 拦截无人机（压制对岸防空）、DWL-200 电子对抗（干扰敌侦察通信）、红旗-17A（掩护渡河车队）、"
       "陆盾-3000（综合防空+对地压制）。")
h3("威胁系数解读")
body("威胁度 0.71、电磁对抗 0.57、时间压力 0.73——高威胁高时间压力，且渡河窗口是“一次性机会”，方案对时序（time）敏感。")
h3("涉及能力维度")
body("火力、机动、防护、电子对抗——渡河是五阶段杀伤链中“突破（breach）”环节最典型的体现。")
h3("实验备注")
body("任务成功率增益 +1.25（65.73→66.98）；值得注意的是论文在渡河场景中 Hope-only 一度反超 Full（−1.03 点），"
     "说明进攻时序类场景中权重调制的作用尤为突出。")

# ---- 场景5 ----
h2("场景五　岛屿补给走廊保障（Island Resupply Corridor）")
h3("任务")
body("在敌方持续袭扰与拦阻下保持岛屿补给走廊畅通——保障（sustain）类场景，比前四个更强调“持续”而非“瞬间达成”。")
h3("地形与天气")
body("海上-岛屿（maritime-island），大风（windy）。海况影响航行与无人机作业。")
h3("兵力对比与敌我装备")
body("我方 4 vs 敌方 2。")
bullet("敌方（说明性编排）：海上快艇袭扰编队、岸基反舰/防空火力、无人机持续监视与拦截。")
bullet("我方（示例编排）：无线电侦测站（海上态势感知）、高功率微波系统（对来袭无人机群软杀伤）、DWL-200 电子对抗"
       "（压制敌通信/导航）、陆盾-3000（护航防空）。")
h3("威胁系数解读")
body("威胁度 0.63、电磁对抗 0.59（全场景最高电磁对抗之一）、时间压力 0.61——“电磁环境差 + 需要长时间保持”的组合，"
     "考验持续保障（sustainment）与电子对抗维度的协同。")
h3("涉及能力维度")
body("持续保障、电子对抗、态势感知、指挥控制——保障类场景维度子集。")
h3("实验备注")
body("任务成功率增益 +2.38（64.05→66.43），在五场景中位居第二，且论文三种子复现中“岛屿场景”是 Reflection-only "
     "少数能反超基线的场景之一。")

# ---- 速览表 ----
h2("五场景一页速览表")
make_table(
    ["场景", "类型", "地形/天气", "核心任务一句话", "兵力(友/敌)", "威胁/EW/时间", "ΔMS"],
    [
        ["沿海突击", "assault", "海岸城市/雨", "夺登陆窗口、压制岸防防空", "5/3", "0.74/0.68/0.77", "+5.15"],
        ["城市防御", "defense", "城市/雾", "守交通通信枢纽、抗装甲突击", "4/2", "0.69/0.54/0.66", "+1.04"],
        ["山地侦察", "recon", "山地/多云", "低能见度下收集火力与补给线情报", "3/2", "0.58/0.46/0.55", "+1.86"],
        ["渡河突破", "assault", "河岸平原/阴", "开辟渡河窗口、巩固桥头堡", "4/2", "0.71/0.57/0.73", "+1.25"],
        ["岛屿补给", "sustain", "海上岛屿/大风", "袭扰拦阻下保持补给走廊畅通", "4/2", "0.63/0.59/0.61", "+2.38"],
    ],
    widths=[2.2, 1.8, 2.4, 5.6, 1.9, 2.6, 1.5],
)
note("ΔMS 为任务成功率增益（Full 相对 Pure LLM），数据来源：论文 Table 7，seed 7、20 次迭代、Full 默认参数；"
     "MS=Mission Success 任务成功率。")

# ============================ 第二部分：仿真原理 ============================
doc.add_page_break()
h1("第二部分　本课题仿真原理与可靠性论证")

h2("一、关键前提：仿真不由大模型执行")
body("一个常见误解是“让大模型在脑子里推演战场”。本系统中，大模型（LLM）的角色被严格限定在“方案生成”："
     "它负责把想定输入组织成结构化方案（分阶段、每阶段含动作类型与动作描述）。方案一旦生成，后续评估完全由独立的"
     "解析计算仿真器完成（代码 engine.py 中的 PlanSimulator 类），LLM 不参与任何战场推演计算。")
body("这一点在架构上被刻意隔离：生成层（LLM）→ 评估层（仿真器），两者通过“结构化方案 JSON”衔接，互不耦合。")

h2("二、仿真原理：能力加权的五阶段杀伤链解析模型")
body("仿真是确定性可计算的解析模型，一次评估分四步完成：")
h3("1. 能力聚合")
body("把我方/敌方每个装备按七维能力（火力 fires、机动 mobility、防护 protection、态势感知 awareness、电子对抗 ew、"
     "持续保障 sustainment、指挥控制 c2）加权聚合，得到蓝方、红方两个能力向量；聚合时考虑装备数量（开根号）与战备度"
     "（对应代码 _aggregate_capabilities）。")
h3("2. 方案→需求映射")
body("把 LLM 生成的方案动作（侦察、电子对抗、打击、机动等动作类型）映射为该阶段的能力需求权重（对应 _phase_requirements，"
     "每种动作类型对应一张能力权重表，如“打击→火力 0.48”）。")
h3("3. 五阶段杀伤链评分")
body("对侦察发现（detect）、压制打断（disrupt）、突破突击（breach）、夺控稳态（control）、持续保障（sustain）逐阶段计算"
     "“蓝方能力×权重 − 红方能力”的差距，经归一化得 0~1 阶段分，并施加地形、天气、平民、电磁环境四项惩罚"
     "（对应 _evaluate_kill_chain）。例如探测阶段公式形如：0.45×感知能力×感知权重 + 0.25×指挥能力×指挥权重 + 0.15×电子战能力×电子战权重 "
     "− 0.30×敌方感知 − 0.18×敌方电子战。")
h3("4. 合成指标与蒙特卡洛")
body("阶段成功率 = 各阶段分×阶段权重；损失、时长、交换比（LER）、综合效能（OE）均由解析公式得出。蒙特卡洛=对阶段分和效能"
     "加 ±3%~5% 的显式受控伪随机扰动（对应代码 rng.uniform(0.97, 1.03)），重复 50 次取均值/置信区间——随机性完全在仿真器内可控，"
     "与 LLM 无关。")
body("一句话概括：这是“能力-需求匹配”的量化推演——方案好不好，取决于它把可用能力对到战场需求上的匹配程度，"
     "而不是大模型的“想象”。", bold=True)

h2("三、与 AFSIM 等三维仿真软件的关系：互补的两层")
make_table(
    ["维度", "本系统（解析仿真）", "AFSIM（三维实体仿真）"],
    [
        ["建模粒度", "能力级（七维能力向量）", "实体级（逐装备弹道/电磁/运动学建模）"],
        ["单次仿真成本", "毫秒级；50 次蒙特卡洛秒级", "分钟~小时级"],
        ["可解释性", "每项得分可分解到“能力×权重×需求”与惩罚项，全链路可追溯", "高保真但“黑箱”，结果难以逐项归因"],
        ["用途", "方案空间快速筛选 + 迭代优化（10-20 轮）", "选定方案高保真验证"],
        ["随机性", "受控 ±3-5% 扰动", "物理随机过程"],
    ],
    widths=[3.0, 6.6, 6.6],
)
body("定位：本系统的仿真不是要替代 AFSIM，而是承担“方案搜索与优化的引擎”角色——AFSIM 单次跑数分钟，无法支撑 "
     "20 轮迭代 × 多候选的搜索空间；解析仿真把这一层成本降了几个数量级。两者是流水线关系：解析仿真快速筛选出有希望的方案"
     "→ 高保真仿真对候选方案做最终验证。这也是论文未来工作中“接入实体级兵棋与外部仿真平台对标”的本意。")

h2("四、可靠性的四条论证")
h3("1. 可解释、可追溯")
body("每个数字都可拆解：阶段分由哪几项能力×权重构成、哪些惩罚项生效、第几轮迭代改变的哪个参数，全部有据可查"
     "（full_result.json 记录每次仿真的阶段明细）。评审可以逐项验算，这是解析模型相对黑箱模型的天然优势。")
h3("2. 受控随机、可复现")
body("蒙特卡洛扰动幅度明确（±3-5%）、随机种子固定（seed 可复现）、50 次聚合给出置信区间。同一方案重跑结果一致"
     "（误差带内），不是“每次都不一样”。")
h3("3. 统计稳健性证据链")
body("多种子复现（5 个种子全为正增益、配对 t 检验 p<0.01）、三个不同大模型后端一致增益、外部基线对照、跨场景 5/5 全胜——"
     "如果仿真不可靠，不可能在所有对照中呈现一致的方向性结论。")
h3("4. 诚实声明边界")
body("论文明确把“仿真未对标外部兵棋平台”列为效度威胁（threats to validity），并声明当前置信区间只反映种内蒙特卡洛方差。"
     "主动说清边界，并指出这正是下一步“接入实体级兵棋/平台对标”的动机，形成逻辑闭环。")

h2("五、一句话总结")
body("“本系统的仿真不是‘大模型推演’，而是一个能力加权的解析计算模型：LLM 只负责把想定写成结构化方案，仿真器把方案映射成"
     "能力需求后做确定性的五阶段杀伤链计算，再用受控蒙特卡洛给出统计量。它的可靠性来自可解释、可复现、可验算，以及多种子/"
     "多模型/跨场景的一致证据；它和 AFSIM 不是竞争关系，而是‘快速迭代筛选 + 高保真验证’的分层流水线。”", bold=True)

note("本文档内容基于代码（military_research_backup/engine.py）与论文（main_blinded.pdf）、用户指导手册 v4.0 整理。")

doc.save("课题背景知识文档_五场景与仿真原理.docx")
print("saved 课题背景知识文档_五场景与仿真原理.docx")
