"""T13：把用户指导手册从 v4.2 更新到 v4.3（研究管线改用 config.json）。

用法： python _t13_apply.py
产出： 反无人机行动方案智能生成系统_用户指导手册_v4.3.docx
"""
from __future__ import annotations

import copy
import datetime
import sys

import docx
from docx.text.paragraph import Paragraph

SRC = "反无人机行动方案智能生成系统_用户指导手册_v4.2.docx"
DST = "反无人机行动方案智能生成系统_用户指导手册_v4.3.docx"
TODAY = datetime.datetime.now().strftime("%Y年%m月%d日")
TODAY_DASH = datetime.datetime.now().strftime("%Y-%m-%d")

doc = docx.Document(SRC)
report: list[str] = []


def set_text(paragraph: Paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._element.getparent().remove(run._element)


def find_para(substring: str, occurrence: int = 1):
    hits = [p for p in doc.paragraphs if substring in p.text]
    if len(hits) < occurrence:
        return None
    return hits[occurrence - 1]


def replace_para(old_substring: str, new_text: str, occurrence: int = 1) -> bool:
    p = find_para(old_substring, occurrence)
    if p is None:
        report.append(f"!! 未找到锚点: {old_substring[:40]}")
        return False
    set_text(p, new_text)
    report.append(f"OK 替换: {old_substring[:40]!r} -> {new_text[:50]!r}")
    return True


def insert_after(anchor_substring: str, text: str, occurrence: int = 1) -> bool:
    anchor = find_para(anchor_substring, occurrence)
    if anchor is None:
        report.append(f"!! 未找到插入锚点: {anchor_substring[:40]}")
        return False
    new_el = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_el)
    new_para = Paragraph(new_el, anchor._parent)
    set_text(new_para, text)
    report.append(f"OK 插入于 {anchor_substring[:32]!r} 之后: {text[:46]!r}")
    return True


def add_table_row_after_header(table, values) -> None:
    row = table.add_row()
    for cell, value in zip(row.cells, values):
        cell.text = value
    tbl = table._tbl
    tbl.remove(row._tr)
    header = table.rows[0]._tr
    header.addnext(row._tr)


# ── 1. 封面版本号 ────────────────────────────────────────────────────────────
replace_para("版本： v4.2", "版本： v4.3")

# ── 2. §5.4 增加模型接入说明 ─────────────────────────────────────────────────
insert_after(
    "系统内置一套完整的军事研究仿真管线",
    "模型接入（v4.3）：研究仿真管线优先读取与前端同一份 config.json"
    "（provider / api_key / model / base_url / max_tokens）；"
    "若未找到可用配置（无 config.json 或 api_key 为空），则回退读取本地模型环境变量 "
    "LOCAL_LLM_BASE_URL、LOCAL_LLM_API_KEY、LOCAL_LLM_MODEL"
    "（默认 http://localhost:1234/v1，兼容 LM Studio）。",
)

# ── 3. §5.6 Tier 1 说明补一句「无需模型」────────────────────────────────────
replace_para(
    "基于案例骨架继承 + HOPE 权重投影 + 兵力威胁比缩放，纯规则出库，不调用大模型。",
    "基于案例骨架继承 + HOPE 权重投影 + 兵力威胁比缩放，纯规则出库，不调用大模型"
    "（无需 API Key，也不需要本地模型服务）。",
)

# ── 4. §6.1 说明 config.json 为管线与前端共用 ─────────────────────────────────
insert_after(
    "配置文件 config.json 格式如下：",
    "该配置文件同时供前端（命令行 / Web / 脚本）与研究仿真管线共用（v4.3 起管线优先读取本文件）。",
)

# ── 5. §6.5 补充管线回退用的本地模型环境变量 ──────────────────────────────────
insert_after(
    "各服务商对应的环境变量名：",
    "研究仿真管线在未找到可用 config.json 时，会回退读取本地模型环境变量："
    "LOCAL_LLM_BASE_URL / LOCAL_LLM_API_KEY / LOCAL_LLM_MODEL"
    "（与上表的云端 API Key 环境变量相互独立）。",
)

# ── 6. 版本历史正文 + 表格 ───────────────────────────────────────────────────
new_history = (
    f"v4.3 — {TODAY} — 研究仿真管线改为优先读取 config.json（与前端共用同一份配置），"
    "未配置时回退 LOCAL_LLM_* 本地模型；重编译 military_research/engine.pyd，"
    "运行清单 full_result.json 的 llm_backend 会如实记录 api_mode（cloud-json / openai-compatible）"
    "与 source。"
)
anchor_v42 = find_para("v4.2 — 2026年08月26日")
if anchor_v42 is None:
    report.append("!! 未找到版本历史 v4.2 条目")
else:
    new_el = copy.deepcopy(anchor_v42._p)
    anchor_v42._p.addprevious(new_el)
    set_text(Paragraph(new_el, anchor_v42._parent), new_history)
    report.append("OK 版本历史新增 v4.3")

replace_para("本手册最后更新：", f"本手册最后更新： {TODAY}")

# ── 7. 附录 版本表格 / 命令速查表 ────────────────────────────────────────────
for table in doc.tables:
    head = " | ".join(c.text.strip() for c in table.rows[0].cells)
    if head.startswith("版本") and "日期" in head:
        add_table_row_after_header(table, ["v4.3", TODAY_DASH, "研究仿真管线改用 config.json（与前端共用），未配置时回退本地模型"])
        report.append("OK 附录版本表新增 v4.3 行")
    if head.startswith("功能") and "命令" in head:
        add_table_row_after_header(table, ["研究仿真管线", "python run_military_research.py --scenario data/sample_antidrone.json --output-dir result/output"])
        report.append("OK 命令速查表新增研究管线行")

doc.save(DST)
print("\n".join(report))
print(f"\n已保存: {DST}")
if any(line.startswith("!!") for line in report):
    sys.exit(1)
